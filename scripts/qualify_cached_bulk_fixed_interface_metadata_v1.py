"""Cached bulk feature axes and old reviewed clinical counts only; no effects."""
from pathlib import Path
import os,sys,json,gzip,hashlib,re,ast
from collections import Counter
from datetime import datetime,timezone
import numpy as np,pandas as pd
from openpyxl import load_workbook

def sha(p):
 h=hashlib.sha256()
 with Path(p).open('rb') as f:
  for b in iter(lambda:f.read(1024**2),b''):h.update(b)
 return h.hexdigest()

def first_column(path,sep='\t'):
 op=gzip.open if path.suffix=='.gz' else open
 with op(path,'rt') as f:
  head=f.readline().rstrip('\n').split(sep)
  axis=[line.rstrip('\n').split(sep,1)[0].strip('"') for line in f]
 return axis,len(head)-1

def field_axis(path):
 if path.suffix=='.xlsx':
  wb=load_workbook(path,read_only=True,data_only=True);ws=wb.worksheets[0]
  axis=[str(r[0]) for r in ws.iter_rows(min_row=2,max_col=1,values_only=True)]
  n=ws.max_column-1;wb.close();return axis,n
 return first_column(path,sep=',' if 'raw_counts.csv' in path.name else '\t')

def main():
 cfg=json.load(sys.stdin)
 base=Path(os.environ['PANCANCER_BULK_ROOT']); old=Path(os.environ['PANCANCER_OLD_BULK_RUN'])
 out=Path(os.environ['PANCANCER_RAM_OUTPUT']);assert out.parent==Path('/dev/shm') and not out.exists();out.mkdir(mode=0o700)
 def save(n,x):(out/n).write_text(json.dumps(x,indent=2,allow_nan=False)+'\n')
 save('private_plan_before_metadata.json',dict(created_UTC=datetime.now(timezone.utc).isoformat(),source_sha256=sha(__file__),fixed_config=cfg,
  scope='All five current complete-expression ICI caches, partial SU2C panel, complete TCGA physical axis. Original expression body scanned only for feature identity; no expression values parsed or effects fitted.',
  no_clinical_patient_row_export=True,no_new_download=True,old_families_unchanged=True))
 geneinfo=base/'ici_bulk/Homo_sapiens.gene_info.gz'; gi=pd.read_csv(geneinfo,sep='\t',usecols=['GeneID','Symbol'],dtype=str)
 assert gi.GeneID.is_unique;id2=gi.set_index('GeneID').Symbol.to_dict()
 probemap=base/'tcga/gencode.v23.annotation.gene.probemap';pm=pd.read_csv(probemap,sep='\t',dtype=str)
 pm['stable']=pm.id.str.split('.').str[0]
 ensgroups=pm.groupby('stable').gene.agg(lambda x:set(x));ens={k:next(iter(v)) for k,v in ensgroups.items() if len(v)==1}
 paths=[('IMvigor210',base/'ici_bulk/imvigor210/data_mrna_seq_tpm.txt','symbol'),
  ('Riaz2017',base/'ici_bulk/riaz2017/GSE91061_raw_counts.csv.gz','Entrez'),
  ('Hugo2016',base/'ici_bulk/hugo2016/GSE78220_FPKM.xlsx','symbol'),
  ('Jung2019',base/'ici_bulk/jung2019_nsclc/GSE135222_exp.tsv.gz','Ensembl'),
  ('Rose2021',base/'ici_bulk/rose2021_urothelial/GSE176307_salmon_tpm.tsv.gz','symbol'),
  ('SU2C_MARK',old/'SU2C_MARK_panel_ready_v1/fixed_panel_TPM_private.tsv.gz','symbol'),
  ('TCGA',base/'tcga/tcga_RSEM_gene_tpm.gz','Ensembl')]
 clinical=[];axes=[]
 for co,path,kind in paths:
  assert path.is_file();axis,n=field_axis(path)
  if kind=='Entrez': symbols=[id2.get(a,'UNMAPPED::'+a) for a in axis]
  elif kind=='Ensembl': symbols=[ens.get(a.split('.')[0],'UNMAPPED::'+a) for a in axis]
  else:symbols=axis
  sets=[];panels={};flat=np.asarray(symbols,str)
  for role,target in enumerate(cfg['interface_symbols']):
   found=[i for i,s in enumerate(symbols) if s==target]
   assert found==np.flatnonzero(flat==target).tolist()
   # Cross-check physical identity with immutable public ids for ID-indexed cohorts.
   if kind=='Entrez':assert found==[i for i,a in enumerate(axis) if a==cfg['interface_Entrez'][role]]
   if kind=='Ensembl':assert found==[i for i,a in enumerate(axis) if a.split('.')[0]==cfg['interface_Ensembl'][role]]
   sets.append(found)
  counter=Counter(symbols)
  for name,genes in cfg['control_panels'].items():
   cards=[counter[g] for g in genes]
   assert cards==[int((flat==g).sum()) for g in genes]
   panels[name]=dict(required=len(genes),unique_physical_columns=sum(c==1 for c in cards),missing_roles=[i for i,c in enumerate(cards) if c==0],ambiguous_roles=[i for i,c in enumerate(cards) if c>1])
  axes.append(dict(cohort=co,cache_scope='fixed_small_panel_only' if co=='SU2C_MARK' else 'complete_cached_expression_input',raw_axis_rows=len(axis),raw_RNA_columns=n,identity_type=kind,
   interface_role_cardinalities=[len(s) for s in sets],all_interface_roles_unique=all(len(s)==1 for s in sets),controls=panels,
   raw_file_bytes=path.stat().st_size,raw_file_SHA256=sha(path),physical_axis_SHA256=hashlib.sha256('\n'.join(axis).encode()).hexdigest(),
   expression_values_not_parsed=True,source_gene_axis_not_proof_expression_or_cell_identity=True))
  # Source-level counts from existing reviewed clinical mappings; individual rows RAM only.
  if co=='TCGA':
   allc=[]
   for p in sorted((old/'analysis_v2').glob('TCGA_*_patient_scores_private.tsv')):
    c=pd.read_csv(p,sep='\t',index_col=0);assert c.index.is_unique
    counts={}
    for ep in ['OS','PFI']:
     e=pd.to_numeric(c[ep],errors='coerce');t=pd.to_numeric(c[ep+'.time'],errors='coerce')
     take=e.isin([0,1]) & t.gt(0) & np.isfinite(t)
     counts[ep]=dict(known_positive_time_units=int(take.sum()),events=int(e[take].sum()))
    allc.append(dict(cancer=p.name.split('_')[1],mapped_primary_tumor_units=len(c),endpoint_counts=counts))
   assert len(allc)==33
   clinical.append(dict(cohort=co,source_cancer_counts=allc,no_outcome_effects_read=True))
  else:
   cp=old/('SU2C_analysis_v1' if co=='SU2C_MARK' else 'analysis_v2')/(co+'_patient_scores_private.tsv')
   c=pd.read_csv(cp,sep='\t',index_col=0);assert c.index.is_unique
   y=pd.to_numeric(c.response,errors='coerce');cnt={}
   for ep in ['OS','PFS']:
    t=pd.to_numeric(c[ep],errors='coerce');e=pd.to_numeric(c[ep+'_event'],errors='coerce')
    take=t.gt(0) & np.isfinite(t) & e.isin([0,1]);cnt[ep]=dict(known_positive_time_units=int(take.sum()),events=int(e[take].sum()))
   clinical.append(dict(cohort=co,old_mapped_units=len(c),response_known=int(y.isin([0,1]).sum()),response_positive=int(y.eq(1).sum()),response_negative=int(y.eq(0).sum()),endpoint_counts=cnt,
    prior_clinical_input_SHA256=sha(cp),no_patient_keys_exported=True,new_question_admission=False))
 save('private_all7_physical_bulk_feature_axis_qualification.json',axes)
 save('private_all6_ICI_and33_TCGA_old_clinical_unit_counts.json',clinical)
 save('private_independent_metadata_proof.json',dict(all_roles_row_vs_vector_and_public_ID_equal=True,all_panels_Counter_vs_vector_equal=True,
  gene_info_SHA256=sha(geneinfo),probemap_SHA256=sha(probemap),patient_unit_truth_inherited_not_newly_proven=True,
  no_expression_outcome_association_or_P=True,no_human_RNA_scores_model_inputs_or_random_arrays_persisted=True))
 files={p.name:dict(bytes=p.stat().st_size,SHA256=sha(p)) for p in out.iterdir() if p.is_file()}
 assert sum(v['bytes'] for v in files.values())<1024**2
 save('private_terminal.json',dict(status='COMPLETE_CACHED_BULK_PHYSICAL_AXES_AND_INHERITED_CLINICAL_COUNTS_NO_EFFECTS',completed_UTC=datetime.now(timezone.utc).isoformat(),source_SHA256=sha(__file__),files=files))
 print(json.dumps(dict(status='COMPLETE',terminal_SHA256=sha(out/'private_terminal.json'))))

if __name__=='__main__':main()
